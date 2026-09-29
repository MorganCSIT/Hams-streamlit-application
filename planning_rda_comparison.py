"""Exploratory planning versus original RDA workbook, independent of PDF filters."""
from io import BytesIO
import math
import re
import unicodedata

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.formatting.rule import ColorScaleRule


def _text(value):
    return "" if pd.isna(value) else str(value).strip()


def _id(value):
    text = re.sub(r"\.0$", "", _text(value))
    return "" if text.lower() in {"nan", "nat", "none", "<na>", "-"} else text


def _ids(values):
    """Accept both scalar IDs and lists already joined by the audit mapping."""
    found = set()
    for value in values:
        for part in re.split(r"[/,;|]", _text(value).strip("()[]")):
            number = _id(part)
            if number:
                found.add(number)
    return found


def _id_list(values):
    ids = _ids(values)
    return "(" + ", ".join(sorted(ids)) + ")" if ids else ""


def client_aliases(frame):
    """Resolve aliases only when the mapping assigns an ID to one main key."""
    frame = frame.copy()
    frame.columns = [str(c).strip().lstrip("\ufeff").lower() for c in frame.columns]
    main = next((c for c in ("client-id", "client_id", "clientid") if c in frame), None)
    if main is None:
        return {}, {}, set()
    number_columns = [c for c in frame if c.startswith("no-client-") or c in {
        "no client", "n° du client", "id client", "client no", "client_nr", "kd-nr", "kd_nr"}]
    aliases, owners = {}, {}
    for row in frame.to_dict("records"):
        key = _id(row.get(main))
        if not key:
            continue
        numbers = _ids(row.get(c) for c in number_columns)
        aliases.setdefault(key, set()).update(numbers)
        for number in numbers:
            owners.setdefault(number, set()).add(key)
    ambiguous = {number for number, keys in owners.items() if len(keys) > 1}
    lookup = {number: next(iter(keys)) for number, keys in owners.items() if len(keys) == 1}
    return lookup, aliases, ambiguous


def _identity(number, name):
    number = re.sub(r"\.0$", "", _text(number))
    if number:
        return "id:" + number
    name = " ".join(unicodedata.normalize("NFKC", _text(name)).casefold().split())
    if not name or set(name.split()) == {"empty"}:
        return ""
    return "name:" + name


def comparison_entries(result):
    names = {}
    collab_aliases = {}
    for row in result.get("map_df", pd.DataFrame()).to_dict("records"):
        cid = _text(row.get("collab_id"))
        names[cid] = (
            _text(row.get("collab_name_sarl")) or _text(row.get("collab_name_wf")))
        collab_aliases.setdefault(cid, set()).update(_ids(row.get(c) for c in (
            "collab_no_sarl", "driverno", "rda_ids", "wf_ids", "planning_ids", "collab_all_ids")))
    client_lookup, client_ids, ambiguous_clients = client_aliases(
        result.get("client_map_df", pd.DataFrame()))
    records = []
    sources = [("Planning", result.get("planning_comparison", result["planning"])),
               ("RDA", result["rda"])]
    for source, frame in sources:
        for pos, row in enumerate(frame.to_dict("records"), 2):
            planning = source == "Planning"
            cid = _text(row.get("collab_id"))
            number = _id(row.get("emp_nr" if planning else "collab_no_sarl"))
            name = _text(row.get("collab_name")) or names.get(cid, "")
            # Unmapped IDs are source-specific: do not invent cross-system matches.
            key = cid or f"unmapped:{source}:{_identity(number, name) or pos}"
            collab_aliases.setdefault(key, set()).update(_ids([number]))
            day = row.get("date" if planning else "jour")
            day = pd.to_datetime(day, errors="coerce")
            day = day.date() if pd.notna(day) else None
            minutes = pd.to_numeric(row.get("duration_min" if planning else "duree_min"), errors="coerce")
            valid_minutes = pd.notna(minutes) and math.isfinite(float(minutes)) and minutes >= 0
            client_name = _text(row.get("client_name"))
            client_id = _id(row.get("client_nr"))
            client_main_id = client_lookup.get(client_id, "")
            client_key = "mapped:" + client_main_id if client_main_id else _identity(client_id, client_name)
            issues = []
            if not cid:
                issues.append("Unmapped collaborator (kept separately)")
            if day is None:
                issues.append("Invalid date (excluded from summaries)")
            if not valid_minutes:
                issues.append("Invalid duration (excluded from minutes)")
            if not client_key and client_name:
                issues.append("Client identity unavailable / anonymized")
            if client_key.startswith("name:"):
                issues.append("Client matched by name only")
            if client_id in ambiguous_clients:
                issues.append("Client ID belongs to multiple mapping keys; alias matching not applied")
            if _text(row.get("_drop")):
                issues.append("PDF planning exclusion: " + _text(row["_drop"]))
            records.append({
                "source": source, "source_row": pos, "date": day,
                "collab_id": key, "collab_number": number, "collab_name": name,
                "client_key": client_key, "client_number": client_id, "client_name": client_name,
                "collab_record_id": number, "client_record_id": client_id,
                "client_main_id": client_main_id,
                "client_all_ids": _id_list(client_ids.get(client_main_id, {client_id})),
                "minutes": float(minutes) if valid_minutes else None,
                "hours": float(minutes) / 60 if valid_minutes else None,
                "client_absent": planning and _text(row.get("client_absent")).upper() == "Y",
                "activity": _text(row.get("event_color" if planning else "prestation_code")),
                "activity_label": _text(row.get("prestation_text")) if not planning else "",
                "planning_type": _text(row.get("type")) if planning else "",
                "issues": "; ".join(issues),
            })
    for row in records:
        row["collab_all_ids"] = _id_list(collab_aliases.get(row["collab_id"], set()))
    return pd.DataFrame(records, columns=[
        "source", "source_row", "date", "collab_id", "collab_number", "collab_name",
        "client_key", "client_number", "client_name", "collab_all_ids", "collab_record_id",
        "client_main_id", "client_all_ids", "client_record_id", "minutes", "hours", "client_absent",
        "activity", "activity_label", "planning_type", "issues"])


def _names(series):
    return " | ".join(sorted({_text(v) for v in series if _text(v)}))


def summarize(entries, keys):
    identity_columns = []
    for entity, key in (("collab", "collab_id"), ("client", "client_key")):
        if key in keys:
            identity_columns += [f"{entity}_all_ids", f"planning_{entity}_record_ids", f"rda_{entity}_record_ids"]
            if entity == "client":
                identity_columns.append("client_main_id")
    columns = list(keys) + ["collab_name", "client_name"] + identity_columns + ["planned_rows", "rda_rows",
        "planned_min", "rda_min", "difference_min", "planned_hours", "rda_hours",
        "difference_hours", "difference_pct", "planned_clients", "rda_clients",
        "matched_clients", "planned_only_clients", "rda_only_clients", "client_counts_match",
        "client_sets_match", "planned_unidentified_rows", "rda_unidentified_rows",
        "invalid_duration_rows", "coverage"]
    records = []
    groups = entries.groupby(list(keys), dropna=False, sort=True) if keys else [((), entries)]
    for key, group in groups:
        if not isinstance(key, tuple):
            key = (key,)
        p = group[group.source.eq("Planning")]
        r = group[group.source.eq("RDA")]
        pc, rc = set(p.client_key) - {""}, set(r.client_key) - {""}
        pm, rm = p.minutes.sum(), r.minutes.sum()
        identities = {}
        for entity, entity_key in (("collab", "collab_id"), ("client", "client_key")):
            if entity_key in keys:
                identities[f"{entity}_all_ids"] = _id_list(group[f"{entity}_all_ids"])
                identities[f"planning_{entity}_record_ids"] = _id_list(p[f"{entity}_record_id"])
                identities[f"rda_{entity}_record_ids"] = _id_list(r[f"{entity}_record_id"])
                if entity == "client":
                    identities["client_main_id"] = _names(group.client_main_id)
        records.append(dict(zip(keys, key)) | {
            **identities,
            "collab_name": _names(group.collab_name) if "collab_id" in keys else "",
            "client_name": _names(group.client_name) if "client_key" in keys else "",
            "planned_rows": len(p), "rda_rows": len(r),
            "planned_min": pm, "rda_min": rm, "difference_min": rm - pm,
            "planned_hours": pm / 60, "rda_hours": rm / 60, "difference_hours": (rm - pm) / 60,
            "difference_pct": (rm - pm) / pm if pm else None,
            "planned_clients": len(pc), "rda_clients": len(rc), "matched_clients": len(pc & rc),
            "planned_only_clients": len(pc - rc), "rda_only_clients": len(rc - pc),
            "client_counts_match": len(pc) == len(rc), "client_sets_match": pc == rc,
            "planned_unidentified_rows": int(p.client_key.eq("").sum()),
            "rda_unidentified_rows": int(r.client_key.eq("").sum()),
            "invalid_duration_rows": int(group.minutes.isna().sum()),
            "coverage": "Both" if len(p) and len(r) else "Planning only" if len(p) else "RDA only",
        })
    return pd.DataFrame(records, columns=columns)


def build_comparison_tables(result):
    entries = comparison_entries(result)
    valid = entries[entries.date.notna()].copy()
    clients = valid[valid.client_key.ne("")]
    active = clients[~clients.client_absent.astype(bool)]
    notes = [
        ("Purpose", "Compare planning with original recorded RDA, before cutting; all uploaded dates, independent of PDF selection."),
        ("Time", "Sum of recorded duration in minutes (hours = minutes / 60), not elapsed shift span. RDA is recorded time, not independently verified work."),
        ("Difference", "RDA minus planning; positive = more recorded than planned. Percentage is blank when planned minutes are zero."),
        ("Overall / Daily / Collaborators", "All activities, including travel and activities without a client. One-sided dates/collaborators remain visible."),
        ("Client_Time", "Only rows with an identifiable client; includes absent-client planning. Active_Client_Time excludes planning rows marked client_absent=Y."),
        ("Client_Detail", "One row per day, collaborator and client; coverage shows planned-only, RDA-only or both. Multiple services are summed, not treated as separate clients."),
        ("Clients", "Distinct clients across the whole group, not the sum of daily counts. Equal counts do not imply equal client sets."),
        ("Identity", "Collaborators use audit mapping; names are displayed. Client aliases use the Matched Clients main key where unambiguous; otherwise client number, with normalized exact name fallback only when no number exists."),
        ("All IDs", "collab_all_ids and client_all_ids list every known source ID linked to that entity, including mapping aliases absent from the current records. Main mapping keys remain separate. No grouping by shared names."),
        ("Record IDs", "collab_record_id and client_record_id are the original IDs on each source entry, not a new row identifier. Summary sheets list planning_*_record_ids and rda_*_record_ids because a summary can contain several source IDs."),
        ("Missing identity", "Blank/empty names without IDs are not counted as distinct clients. Unidentified row counts are shown; client match flags cover identifiable clients only. ID-only and name-only entries are not automatically linked."),
        ("Names", "Different IDs remain separate even when every name is empty. Unmapped collaborators remain source-specific and appear in Data_Quality."),
        ("Dates and quality", "Invalid dates are excluded from summaries but retained in Entries/Data_Quality. Invalid or negative durations are excluded from minutes and counted. Missing-source time is zero, not proof of no work."),
        ("Activities", "Planning colors/types and RDA service codes are separate taxonomies, shown separately without assuming equivalence. No automatic exclusion of breaks or absences from total time."),
        ("Absences", "Absences shows planned client_absent=Y entries. Active_Client_Time still includes all client-linked RDA; use Client_Detail to investigate."),
        ("Duplicates / overlaps", "Rows are not deduplicated; overlapping service durations can add up beyond clock time. Row counts represent entries, not visits."),
        ("Source rows", "Entries source_row is the data row + header in the selected source sheet. Personal contact details and free-text care notes are not exported."),
    ]
    daily = summarize(valid, ["date", "collab_id"])
    tables = {
        "Read_Me": pd.DataFrame(notes, columns=["Topic", "Definition"]),
        "Overall": summarize(valid, []),
        "Daily": summarize(valid, ["date"]),
        "Collaborators": summarize(valid, ["collab_id"]),
        "Day_Collaborator": daily,
        "Client_Time": summarize(clients, ["date", "collab_id"]),
        "Active_Client_Time": summarize(active, ["date", "collab_id"]),
        "Clients": summarize(clients, ["client_key"]),
        "Client_Detail": summarize(clients, ["date", "collab_id", "client_key"]),
        "Exceptions": daily[(daily.difference_min.abs() > 0.01) | ~daily.client_sets_match |
                            daily.coverage.ne("Both") | daily.invalid_duration_rows.gt(0)].copy(),
        "Activities": valid.groupby(["source", "activity", "activity_label", "planning_type"], dropna=False).agg(
            rows=("source", "size"), minutes=("minutes", "sum"), hours=("hours", "sum")).reset_index(),
        "Absences": entries[entries.client_absent.astype(bool)].copy(),
        "Data_Quality": entries[entries.issues.ne("")].copy(),
        "Entries": entries,
    }
    return tables


def build_comparison_workbook(result):
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for title, frame in build_comparison_tables(result).items():
            frame.to_excel(writer, sheet_name=title, index=False)
            sheet = writer.sheets[title]
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(color="FFFFFF", bold=True)
                cell.fill = PatternFill("solid", fgColor="234E70")
            for column in sheet.columns:
                letter = column[0].column_letter
                header = column[0].value
                sheet.column_dimensions[letter].width = min(48, max(14, len(str(header)) + 2))
                for cell in column[1:]:
                    # Names remain literal text even when beginning with '='.
                    if cell.data_type == "f":
                        cell.data_type = "s"
                    if isinstance(cell.value, (float, int)) and not isinstance(cell.value, bool):
                        cell.number_format = "0.00%" if header == "difference_pct" else "0.00" if any(
                            word in header for word in ("min", "hours")) else "0"
                if header == "difference_min" and sheet.max_row > 1:
                    sheet.conditional_formatting.add(f"{letter}2:{letter}{sheet.max_row}",
                        ColorScaleRule(start_type="min", start_color="63BE7B", mid_type="num",
                                       mid_value=0, mid_color="FFFFFF", end_type="max", end_color="F8696B"))
            if title == "Read_Me":
                from openpyxl.styles import Alignment
                sheet.column_dimensions["B"].width = 110
                for row in sheet.iter_rows(min_row=2):
                    row[1].alignment = Alignment(wrap_text=True, vertical="top")
                    sheet.row_dimensions[row[0].row].height = 45
    output.seek(0)
    return output

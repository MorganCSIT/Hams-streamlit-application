import unittest
from io import BytesIO
from unittest.mock import patch
from tempfile import TemporaryDirectory
from pathlib import Path
from zipfile import ZipFile

import pandas as pd
from openpyxl import load_workbook

from planning_rda_comparison import build_comparison_tables, build_comparison_workbook


def fixture():
    return {
        "planning": pd.DataFrame([
            dict(collab_id="a", date="2026-08-01", duration_min=60, client_nr="1", client_name="empty", client_absent="N"),
            dict(collab_id="a", date="2026-08-01", duration_min=30, client_nr="1", client_name="empty", client_absent="Y"),
            dict(collab_id="b", date="2026-08-02", duration_min=20, client_nr="3", client_name="empty"),
        ]),
        "rda": pd.DataFrame([
            dict(collab_id="a", jour="2026-08-01", duree_min=45, client_nr="2", client_name="empty"),
            dict(collab_id="a", jour="2026-08-01", duree_min=15, client_nr=None, prestation_code="61010"),
            dict(collab_id="c", jour="2026-08-03", duree_min=10, client_nr="3", client_name="empty"),
        ]),
        "map_df": pd.DataFrame([dict(collab_id="a", collab_name_sarl="=literal name")]),
    }


class ComparisonTests(unittest.TestCase):
    def test_alias_lists_and_original_record_ids(self):
        result = fixture()
        result["map_df"] = pd.DataFrame([
            dict(collab_id="a", rda_ids="2343/3432", wf_ids="4322", planning_ids="43432"),
            dict(collab_id="a", rda_ids="2343", wf_ids="9999"),
        ])
        result["planning"]["emp_nr"] = [43432, 2343, 77]
        result["rda"]["collab_no_sarl"] = [3432, 3432, 88]
        result["client_map_df"] = pd.DataFrame([
            {"client-id": "client-a", "no-client-sa-101": 1, "no-client-sarl-102": 2},
            {"client-id": "client-a", "no-client-sa-101": 1, "no-client-sarl-102": 99},
        ])
        tables = build_comparison_tables(result)
        entries = tables["Entries"]
        self.assertEqual(entries.iloc[0].collab_all_ids, "(2343, 3432, 4322, 43432, 9999)")
        self.assertEqual(entries.iloc[0].collab_record_id, "43432")
        self.assertEqual(entries.iloc[3].collab_record_id, "3432")
        self.assertEqual(entries.iloc[0].client_record_id, "1")
        self.assertEqual(entries.iloc[3].client_record_id, "2")
        self.assertEqual(entries.iloc[0].client_all_ids, "(1, 2, 99)")
        self.assertEqual(entries.iloc[0].client_main_id, "client-a")
        daily = tables["Day_Collaborator"].set_index("collab_id")
        self.assertTrue(daily.loc["a", "client_sets_match"])
        self.assertEqual(daily.loc["a", "planning_collab_record_ids"], "(2343, 43432)")
        self.assertEqual(daily.loc["a", "rda_collab_record_ids"], "(3432)")
        detail = tables["Client_Detail"]
        matched = detail[detail.client_main_id.eq("client-a")].iloc[0]
        self.assertEqual(matched.coverage, "Both")
        self.assertEqual(matched.planning_client_record_ids, "(1)")
        self.assertEqual(matched.rda_client_record_ids, "(2)")
        self.assertEqual(tables["Overall"].iloc[0].planned_min, 110)

    def test_ambiguous_client_mapping_does_not_merge_main_keys(self):
        result = fixture()
        result["client_map_df"] = pd.DataFrame([
            {"client-id": "x", "no-client-sa-101": 1},
            {"client-id": "y", "no-client-sa-101": 1},
        ])
        tables = build_comparison_tables(result)
        first = tables["Entries"].iloc[0]
        self.assertEqual(first.client_key, "id:1")
        self.assertEqual(first.client_main_id, "")
        self.assertEqual(first.client_all_ids, "(1)")
        self.assertIn("multiple mapping keys", first.issues)

    def test_totals_counts_and_one_sided_groups(self):
        tables = build_comparison_tables(fixture())
        total = tables["Overall"].iloc[0]
        self.assertEqual((total.planned_min, total.rda_min, total.difference_min), (110, 70, -40))
        daily = tables["Day_Collaborator"].set_index("collab_id")
        self.assertTrue(daily.loc["a", "client_counts_match"])
        self.assertFalse(daily.loc["a", "client_sets_match"])
        self.assertEqual(daily.loc["a", "planned_clients"], 1)
        self.assertEqual(daily.loc["b", "coverage"], "Planning only")
        self.assertEqual(daily.loc["c", "coverage"], "RDA only")
        self.assertEqual(tables["Client_Time"].iloc[0].rda_min, 45)
        self.assertEqual(tables["Active_Client_Time"].iloc[0].planned_min, 60)
        self.assertEqual(len(tables["Absences"]), 1)

    def test_invalid_rows_and_unmapped_identities_are_visible(self):
        result = fixture()
        result["planning_comparison"] = pd.DataFrame([
            dict(collab_id=None, emp_nr="9", date="2026-08-01", noted_duration_min=-5, duration_min=-5, client_name="empty empty"),
            dict(collab_id="a", date="bad", duration_min=100),
        ])
        tables = build_comparison_tables(result)
        self.assertEqual(tables["Overall"].iloc[0].planned_min, 0)
        self.assertEqual(tables["Overall"].iloc[0].invalid_duration_rows, 1)
        self.assertEqual(len(tables["Data_Quality"]), 2)
        self.assertEqual(len(tables["Entries"]), 5)
        self.assertEqual(tables["Overall"].iloc[0].planned_clients, 0)

    def test_empty_inputs_and_literal_names(self):
        empty = {"planning": pd.DataFrame(), "rda": pd.DataFrame()}
        self.assertEqual(build_comparison_tables(empty)["Overall"].iloc[0].planned_min, 0)
        workbook = load_workbook(build_comparison_workbook(fixture()))
        self.assertEqual(len(workbook.sheetnames), 15)
        sheet = workbook["Collaborators"]
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet["B2"].value, "=literal name")
        self.assertEqual(sheet["B2"].data_type, "s")

    def test_duration_reconciliation_and_mapped_scope(self):
        result = fixture()
        result["planning_comparison"] = pd.DataFrame([
            dict(collab_id="a", date="2026-08-01", noted_duration_min=20, duration_min=20,
                 start="2026-08-01 10:00", end="2026-08-01 10:30"),
            dict(collab_id="a", date="2026-08-01", noted_duration_min=None, duration_min=30,
                 start="2026-08-01 10:00", end="2026-08-01 10:30"),
            dict(collab_id=None, emp_nr="99", date="2026-08-01", noted_duration_min=10, duration_min=10),
            dict(collab_id="a", date="2026-08-01", noted_duration_min=-5, duration_min=-5,
                 start="2026-08-01 10:00", end="2026-08-01 10:30"),
        ])
        tables = build_comparison_tables(result)
        total = tables["Overall"].iloc[0]
        self.assertEqual(total.planned_min, 60)
        self.assertEqual(total.planned_noted_min, 30)
        self.assertEqual(total.planned_fallback_min, 30)
        self.assertEqual(total.planned_fallback_rows, 1)
        self.assertEqual(total.duration_mismatch_rows, 1)
        self.assertEqual(total.invalid_duration_rows, 1)
        self.assertEqual(len(tables["Unassigned"]), 1)
        self.assertEqual(tables["Collaborators"].planned_min.sum(), 50)
        self.assertEqual(tables["Unassigned"].minutes.sum(), 10)
        self.assertEqual(len(build_comparison_tables(result, ("2026-08-02", "2026-08-03"))["Entries"]), 1)

    def test_reference_workbook_schema(self):
        reference = Path(__file__).resolve().parents[1] / "planning_rda_comparison.xlsx"
        if not reference.exists():
            self.skipTest("Uploaded reference workbook unavailable")
        workbook = load_workbook(reference, read_only=True)
        tables = build_comparison_tables(fixture())
        self.assertEqual(list(tables), workbook.sheetnames)
        for name, table in tables.items():
            self.assertEqual(list(table.columns), list(next(workbook[name].values)), name)
        workbook.close()

    def test_complete_package_includes_new_sibling_folder(self):
        import audit_webfleet_rda as audit
        result = fixture()
        result["excel_bytes"] = BytesIO(b"existing audit")
        with TemporaryDirectory() as tmp, \
                patch.object(audit, "get_session_output_root", return_value=Path(tmp)), \
                patch.object(audit, "audit_run_rda_cutting", return_value={}), \
                patch.object(audit, "audit_generate_pdfs", return_value=None), \
                patch.object(audit, "audit_build_rda_cutting_package", return_value={}):
            package = audit.audit_build_complete_package(result)
            with ZipFile(package["zip_bytes"]) as archive:
                self.assertIn("Webfleet_RDA_Audit_report/audit_report.xlsx", archive.namelist())
                data = archive.read("Planning_RDA_Comparison/planning_rda_comparison.xlsx")
                self.assertIn("Client_Detail", load_workbook(BytesIO(data)).sheetnames)


if __name__ == "__main__":
    unittest.main()
